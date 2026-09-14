from docx import Document
from pathlib import Path

src = Path(r"C:/Users/Administrator/Desktop") / "AI学习" / "未来Agent项目.docx"
d = Document(str(src))
lines = []
lines.append("PARAS " + str(len(d.paragraphs)))
lines.append("TABLES " + str(len(d.tables)))
lines.append("---PARAS---")
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t:
        lines.append("[" + str(i) + "] " + t)
lines.append("---TABLES---")
for ti, table in enumerate(d.tables):
    lines.append("===TABLE " + str(ti) + " r=" + str(len(table.rows)) + " c=" + str(len(table.columns)) + "===")
    for ri, row in enumerate(table.rows):
        cells = [c.text.strip().replace("\n", " | ") for c in row.cells]
        lines.append("R" + str(ri) + ": " + " || ".join(cells))
out = Path(r"D:/AI.project/AI-assiant/.codex_tmp/agent_doc_utf8.txt")
out.write_text("\n".join(lines), encoding="utf-8")
print(out)
