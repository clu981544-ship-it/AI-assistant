from docx import Document

path = r"C:\Users\Administrator\Desktop\AI学习\未来Agent项目.docx"
d = Document(path)
print("PARAS", len(d.paragraphs))
print("TABLES", len(d.tables))
print("---PARAS---")
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t:
        print(f"[{i}] {t}")
print("---TABLES---")
for ti, table in enumerate(d.tables):
    print(f"===TABLE {ti} r={len(table.rows)} c={len(table.columns)}===")
    for ri, row in enumerate(table.rows):
        cells = [c.text.strip().replace("\n", " | ") for c in row.cells]
        print(f"R{ri}: " + " || ".join(cells))
